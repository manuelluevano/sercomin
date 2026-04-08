#!/usr/bin/env python3

import json
import re
import sys
import tempfile
import zipfile
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parent.parent
TEMPLATE_PATH = PROJECT_ROOT / "app" / "assets" / "plantilla" / "PLANTILLA.xlsx"
BASE_ITEM_START_ROW = 14
TOTAL_ROW_OFFSET = 1
IVA_ROW_OFFSET = 2
GRAND_TOTAL_ROW_OFFSET = 3
NOTES_START_OFFSET = 6
WORKER_ROW_OFFSET = 9


def normalize_text(value: str) -> str:
    return value or ""


def patch_sheet_xml(template_xml: str, generated_xml: str) -> bytes:
    template_open = re.search(r"(<worksheet[^>]*>)", template_xml)
    generated_open = re.search(r"(<worksheet[^>]*>)", generated_xml)
    template_tail = re.search(r"(<pageMargins[\s\S]*</worksheet>)", template_xml)
    generated_tail = re.search(r"(<pageMargins[\s\S]*</worksheet>)", generated_xml)

    if not template_open or not generated_open or not template_tail or not generated_tail:
        raise RuntimeError("No se pudo reconstruir la hoja principal de la plantilla.")

    patched_xml = generated_xml
    patched_xml = patched_xml.replace(generated_open.group(1), template_open.group(1), 1)
    patched_xml = patched_xml.replace(generated_tail.group(1), template_tail.group(1), 1)

    if template_xml.startswith("<?xml") and not patched_xml.startswith("<?xml"):
        declaration = template_xml.split("\n", 1)[0]
        patched_xml = f"{declaration}\n{patched_xml}"

    return patched_xml.encode("utf-8")


def clone_item_row(sheet, source_row: int, target_row: int) -> None:
    source_dimensions = sheet.row_dimensions[source_row]
    target_dimensions = sheet.row_dimensions[target_row]
    target_dimensions.height = source_dimensions.height
    target_dimensions.hidden = source_dimensions.hidden

    for column in range(1, 13):
        source_cell = sheet.cell(source_row, column)
        target_cell = sheet.cell(target_row, column)
        if source_cell.has_style:
            target_cell._style = copy(source_cell._style)
        target_cell.font = copy(source_cell.font)
        target_cell.fill = copy(source_cell.fill)
        target_cell.border = copy(source_cell.border)
        target_cell.alignment = copy(source_cell.alignment)
        target_cell.protection = copy(source_cell.protection)
        target_cell.number_format = source_cell.number_format
        target_cell.value = None

    sheet.merge_cells(start_row=target_row, start_column=4, end_row=target_row, end_column=5)


def fill_workbook(payload: dict, output_path: Path) -> None:
    workbook = load_workbook(TEMPLATE_PATH)
    sheet = workbook[workbook.sheetnames[0]]

    items = payload["items"]
    if not items:
        raise RuntimeError("Agrega al menos una partida antes de generar el Excel.")

    extra_rows = max(len(items) - 1, 0)
    if extra_rows:
        sheet.insert_rows(BASE_ITEM_START_ROW + 1, amount=extra_rows)
        for target_row in range(BASE_ITEM_START_ROW + 1, BASE_ITEM_START_ROW + 1 + extra_rows):
            clone_item_row(sheet, BASE_ITEM_START_ROW, target_row)

    sheet["J3"] = payload["folio"]
    sheet["J10"] = datetime.strptime(payload["quoteDate"], "%Y-%m-%d")
    sheet["A11"] = f"Con At´n:  {normalize_text(payload['clientName']) or 'NOMBRE CLIENTE'}"

    notes_start_row = BASE_ITEM_START_ROW + NOTES_START_OFFSET + extra_rows
    worker_row = BASE_ITEM_START_ROW + WORKER_ROW_OFFSET + extra_rows
    notes = payload.get("notes", [])
    defaults = [
        "** Precios en PESOS",
        "** Precios más el 16% de IVA",
        "** LAB: Guadalajara",
        "** Tiempo de entrega 1/2 DIA",
        "**Precios y existencia sujeta a cambios.",
        f"** Vigencia de cotización {payload['validity']}",
        "** Sujeto a disponibilidad",
        "**Penalizacion por cancelacion del 40%",
    ]
    for idx in range(8):
        sheet[f"A{notes_start_row + idx}"] = notes[idx] if idx < len(notes) else defaults[idx]

    sheet[f"J{worker_row}"] = (
        normalize_text(payload.get("workerName", "")) or "Ing. TRABAJADOR QUE INGRESOA. COTIZADOR"
    )

    for offset, item in enumerate(items):
        row = BASE_ITEM_START_ROW + offset
        sheet[f"A{row}"] = item["cantidad"]
        sheet[f"B{row}"] = normalize_text(item["unidad"])
        sheet[f"D{row}"] = normalize_text(item["descripcion"])
        sheet[f"F{row}"] = normalize_text(item["numeroParte"]) or None
        sheet[f"G{row}"] = item["existencia"]
        sheet[f"H{row}"] = normalize_text(item["entrega"])
        sheet[f"I{row}"] = item["precioUnitario"]
        sheet[f"J{row}"] = f"=I{row}*A{row}"

    total_row = BASE_ITEM_START_ROW + TOTAL_ROW_OFFSET + extra_rows
    iva_row = BASE_ITEM_START_ROW + IVA_ROW_OFFSET + extra_rows
    grand_total_row = BASE_ITEM_START_ROW + GRAND_TOTAL_ROW_OFFSET + extra_rows
    last_item_row = BASE_ITEM_START_ROW + len(items) - 1

    sheet[f"I{total_row}"] = "SUMA"
    sheet[f"J{total_row}"] = f"=SUM(J{BASE_ITEM_START_ROW}:J{last_item_row})"
    sheet[f"I{iva_row}"] = "IVA 16%"
    sheet[f"J{iva_row}"] = f"=J{total_row}*0.16"
    sheet[f"I{grand_total_row}"] = "TOTAL"
    sheet[f"J{grand_total_row}"] = f"=J{total_row}+J{iva_row}"

    workbook.save(output_path)


def main():
    payload_path = Path(sys.argv[1])
    output_dir = Path(sys.argv[2])
    payload = json.loads(payload_path.read_text(encoding="utf-8"))

    output_dir.mkdir(parents=True, exist_ok=True)
    safe_folio = payload["folio"].replace(" ", "_").replace("/", "-")
    xlsx_path = output_dir / f"{safe_folio}.xlsx"

    with tempfile.TemporaryDirectory(prefix="sercomin-xlsx-") as temp_dir:
        generated_path = Path(temp_dir) / "generated.xlsx"
        fill_workbook(payload, generated_path)

        with zipfile.ZipFile(TEMPLATE_PATH, "r") as template_zip:
            template_files = {name: template_zip.read(name) for name in template_zip.namelist()}

        with zipfile.ZipFile(generated_path, "r") as generated_zip:
            generated_files = {name: generated_zip.read(name) for name in generated_zip.namelist()}
            generated_sheet_xml = generated_zip.read("xl/worksheets/sheet1.xml").decode("utf-8")

        template_sheet_xml = template_files["xl/worksheets/sheet1.xml"].decode("utf-8")
        generated_files["xl/worksheets/sheet1.xml"] = patch_sheet_xml(
            template_sheet_xml,
            generated_sheet_xml,
        )

        for name, content in template_files.items():
            if (
                name == "[Content_Types].xml"
                or name == "xl/sharedStrings.xml"
                or name == "xl/worksheets/_rels/sheet1.xml.rels"
                or name.startswith("xl/drawings/")
                or name.startswith("xl/media/")
                or name.startswith("xl/printerSettings/")
            ):
                generated_files[name] = content

        with zipfile.ZipFile(xlsx_path, "w", compression=zipfile.ZIP_DEFLATED) as out_zip:
            for name, content in generated_files.items():
                out_zip.writestr(name, content)

    print(json.dumps({"xlsx_path": str(xlsx_path), "file_name": xlsx_path.name}))


if __name__ == "__main__":
    main()
